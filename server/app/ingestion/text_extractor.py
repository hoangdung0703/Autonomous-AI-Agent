"""Text extraction from source documents (Section 6.1 of requirements.md).

PDF  -> pdfplumber first; if the extracted text is shorter than
        MIN_TEXT_CHARS, the file is treated as scanned and ScannedDocumentError
        is raised so document_ingester can route it to ocr_extractor instead.
DOCX -> python-docx paragraph extraction.
XLSX -> openpyxl, each sheet converted to a CSV-like text block.
"""

import re
from pathlib import Path

import pdfplumber
from docx import Document
from openpyxl import load_workbook

MIN_TEXT_CHARS = 100

# pdfplumber emits "(cid:N)" placeholders when a PDF font's characters don't
# map cleanly to Unicode (seen in practice on some bullet/symbol glyphs) —
# these are extraction artifacts, never real content, so they're stripped
# before the text reaches the chunker.
_CID_ARTIFACT_RE = re.compile(r"\(cid:\d+\)")


class ScannedDocumentError(Exception):
    """Raised when a PDF's extracted text is too short to be real content —
    signals document_ingester to fall back to Gemini Vision OCR."""


async def extract_text(file_path: str) -> str:
    suffix = Path(file_path).suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(file_path)
    if suffix == ".docx":
        return _extract_docx(file_path)
    if suffix == ".xlsx":
        return _extract_xlsx(file_path)
    raise ValueError(f"Unsupported file type: {suffix}")


def _extract_pdf(file_path: str) -> str:
    with pdfplumber.open(file_path) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    text = _CID_ARTIFACT_RE.sub("", text)
    if len(text.strip()) < MIN_TEXT_CHARS:
        raise ScannedDocumentError(
            f"{file_path} has fewer than {MIN_TEXT_CHARS} extractable characters — likely a scanned document"
        )
    return text


def _extract_docx(file_path: str) -> str:
    document = Document(file_path)
    return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())


def _extract_xlsx(file_path: str) -> str:
    workbook = load_workbook(file_path, data_only=True)
    blocks = []
    for sheet in workbook.worksheets:
        rows = [
            ",".join("" if cell is None else str(cell) for cell in row)
            for row in sheet.iter_rows(values_only=True)
        ]
        blocks.append(f"Sheet: {sheet.title}\n" + "\n".join(rows))
    return "\n\n".join(blocks)
